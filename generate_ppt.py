#!/usr/bin/env python3
"""
generate_ppt.py - Generates a 16:9 PowerPoint presentation for the
Linux Virtual Character Device Driver project.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path="virtual_char_device_driver.pptx"):
    prs = Presentation()
    # Set slide dimensions to 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette: Modern Dark Tech Theme
    BG_DARK = RGBColor(18, 24, 38)       # #121826 Navy Slate
    CARD_BG = RGBColor(28, 38, 58)       # #1C263A Lighter card slate
    ACCENT_CYAN = RGBColor(34, 211, 238) # #22D3EE Cyan highlight
    ACCENT_BLUE = RGBColor(59, 130, 246) # #3B82F6 Vibrant blue
    TEXT_WHITE = RGBColor(248, 250, 252) # #F8FAFC Pure white text
    TEXT_MUTED = RGBColor(148, 163, 184) # #94A3B8 Muted text
    ACCENT_GREEN = RGBColor(74, 222, 128)# #4ADE80 Green success
    ACCENT_ORANGE = RGBColor(251, 146, 60)# #FB923C Orange accent

    def set_slide_background(slide, color=BG_DARK):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="LINUX KERNEL DEVICE DRIVER"):
        # Category tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category_text.upper()
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = ACCENT_CYAN
        p_c.font.name = "Arial"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE
        p_t.font.name = "Arial"

    def add_card(slide, left, top, width, height, title=None, border_color=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.fill.background()

        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.4))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = ACCENT_CYAN
            p.font.name = "Arial"
        return card

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, RGBColor(12, 16, 28))

    # Center card
    add_card(slide1, Inches(1.5), Inches(1.2), Inches(10.33), Inches(5.1), border_color=ACCENT_BLUE)

    title_box = slide1.shapes.add_textbox(Inches(2.0), Inches(1.8), Inches(9.33), Inches(2.2))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "SYSTEMS PROGRAMMING & OS ARCHITECTURE"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_CYAN
    p0.font.name = "Arial"

    p1 = tf1.add_paragraph()
    p1.text = "Linux Virtual Character Device Driver"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.font.name = "Arial"
    p1.space_before = Pt(10)

    p2 = tf1.add_paragraph()
    p2.text = "A Complete Kernel Module & C++ User-Space Interface for Ubuntu 24.04"
    p2.font.size = Pt(18)
    p2.font.color.rgb = ACCENT_GREEN
    p2.font.name = "Arial"
    p2.space_before = Pt(8)

    meta_box = slide1.shapes.add_textbox(Inches(2.0), Inches(4.3), Inches(9.33), Inches(1.6))
    tf_m = meta_box.text_frame
    tf_m.word_wrap = True

    pm1 = tf_m.paragraphs[0]
    pm1.text = "• Implementation: C (Kernel Module `mydev.c`) + C++17 (`app.cpp`) + Kbuild Makefile"
    pm1.font.size = Pt(13)
    pm1.font.color.rgb = TEXT_MUTED

    pm2 = tf_m.add_paragraph()
    pm2.text = "• Core Concepts: System Calls, VFS, ioctl, Mutex Concurrency, Ring 0 vs Ring 3 Isolation"
    pm2.font.size = Pt(13)
    pm2.font.color.rgb = TEXT_MUTED
    pm2.space_before = Pt(4)

    # =========================================================================
    # SLIDE 2: Project Overview & Objectives
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "Project Overview & Key Objectives")

    # Card 1: What is a Virtual Character Device?
    add_card(slide2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), title="Virtual Character Device Concept", border_color=ACCENT_CYAN)
    tb = slide2.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True
    bullets = [
        ("Byte-Stream Interface: ", "Unlike block devices (disks), character devices are accessed sequentially as a stream of bytes."),
        ("No Physical Hardware Required: ", "The backing hardware is simulated entirely using a 1 KB buffer in kernel RAM."),
        ("Standard File Abstraction: ", "Programs interact with the driver using standard POSIX functions: open(), read(), write(), lseek(), ioctl(), close()."),
        ("Dynamic Device Node: ", "Registered as `/dev/mydev` with dynamic major/minor allocation and automatic udev permissions.")
    ]
    for i, (b_title, b_desc) in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(10)
        run_b = p.add_run()
        run_b.text = "• " + b_title
        run_b.font.bold = True
        run_b.font.size = Pt(13)
        run_b.font.color.rgb = ACCENT_CYAN
        run_d = p.add_run()
        run_d.text = b_desc
        run_d.font.size = Pt(13)
        run_d.font.color.rgb = TEXT_WHITE

    # Card 2: Key Features Implemented
    add_card(slide2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), title="Deliverables & Capabilities", border_color=ACCENT_BLUE)
    tb2 = slide2.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.3))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    bullets2 = [
        ("Kernel Module (`mydev.c`): ", "Registers `/dev/mydev`, handles VFS calls, mutex protection, and kernel logging with printk()."),
        ("Shared Interface (`mydev_ioctl.h`): ", "Defines ioctl command numbers and buffer constants shared across user & kernel space."),
        ("C++17 User Application (`app.cpp`): ", "Modern interactive menu utilizing RAII class wrappers, system_error exceptions, and robust I/O."),
        ("Two-Pass Makefile: ", "Coordinates Kbuild for module compilation and g++ for the user-space client application."),
        ("Lifecycle Script (`mydev.sh`): ", "Automates module loading, unloading, udev rule installation (MODE 0666), and status reporting.")
    ]
    for i, (b_title, b_desc) in enumerate(bullets2):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.space_before = Pt(8)
        run_b = p.add_run()
        run_b.text = "• " + b_title
        run_b.font.bold = True
        run_b.font.size = Pt(13)
        run_b.font.color.rgb = ACCENT_GREEN
        run_d = p.add_run()
        run_d.text = b_desc
        run_d.font.size = Pt(13)
        run_d.font.color.rgb = TEXT_WHITE

    # =========================================================================
    # SLIDE 3: Architecture & System Flow
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "System Architecture: User Space to Kernel Space")

    # 3 Horizontal Layers / Cards
    add_card(slide3, Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.5), title="1. USER SPACE (Ring 3 - Unprivileged)", border_color=ACCENT_BLUE)
    tb_u = slide3.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(11.3), Inches(0.9))
    tf_u = tb_u.text_frame
    tf_u.word_wrap = True
    p = tf_u.paragraphs[0]
    p.text = "• C++ User App (`./app`) / Shell Tools (`cat`, `echo`) -> Invokes Standard POSIX System Calls: open(), read(), write(), ioctl(), close()"
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_WHITE
    p2 = tf_u.add_paragraph()
    p2.text = "• File Descriptor (fd): A process-local handle (e.g., fd=3) indexing the process's open file table."
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(4)

    add_card(slide3, Inches(0.8), Inches(3.3), Inches(11.7), Inches(1.5), title="2. VFS & DEVICE NODE ABSTRACTION (/dev/mydev)", border_color=ACCENT_CYAN)
    tb_v = slide3.shapes.add_textbox(Inches(1.0), Inches(3.8), Inches(11.3), Inches(0.9))
    tf_v = tb_v.text_frame
    tf_v.word_wrap = True
    p = tf_v.paragraphs[0]
    p.text = "• Inode & Device Node: `/dev/mydev` stores (Major Number = 240 [Driver ID], Minor Number = 0 [Instance ID])."
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_WHITE
    p2 = tf_v.add_paragraph()
    p2.text = "• Virtual File System (VFS): Translates generic read/write calls into `struct file_operations mydev_fops` function pointers."
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(4)

    add_card(slide3, Inches(0.8), Inches(5.0), Inches(11.7), Inches(1.8), title="3. KERNEL SPACE (Ring 0 - Privileged Driver `mydev.ko`)", border_color=ACCENT_GREEN)
    tb_k = slide3.shapes.add_textbox(Inches(1.0), Inches(5.5), Inches(11.3), Inches(1.2))
    tf_k = tb_k.text_frame
    tf_k.word_wrap = True
    p = tf_k.paragraphs[0]
    p.text = "• Handlers: `mydev_read()`, `mydev_write()`, `mydev_ioctl()`, `mydev_open()`, `mydev_release()`."
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_WHITE
    p2 = tf_k.add_paragraph()
    p2.text = "• Concurrency & Memory: Protected by `mutex mydev_lock`. Safe copies via `copy_from_user()` / `copy_to_user()`."
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(4)
    p3 = tf_k.add_paragraph()
    p3.text = "• Backing Storage: `device_buffer[1024]` (1 KB Kernel RAM) serving as the simulated physical hardware."
    p3.font.size = Pt(12)
    p3.font.color.rgb = ACCENT_GREEN
    p3.space_before = Pt(4)

    # =========================================================================
    # SLIDE 4: Kernel Module Implementation Details
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "Kernel Module Architecture (`mydev.c`)")

    add_card(slide4, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2), title="File Operations Table (fops)", border_color=ACCENT_CYAN)
    tb = slide4.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.3), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True
    items = [
        ("mydev_open()", "Tracks active handles with atomic_inc_return(&open_count) and logs to dmesg."),
        ("mydev_release()", "Decrements open handle count upon close()."),
        ("mydev_write()", "Locks mutex, validates offset, copies user bytes to `device_buffer` via copy_from_user(), updates length."),
        ("mydev_read()", "Locks mutex, returns EOF (0) when past length, copies to user via copy_to_user(), advances file position."),
        ("mydev_llseek()", "Supports SEEK_SET, SEEK_CUR, SEEK_END within 0..1024 bound."),
        ("mydev_ioctl()", "Dispatches custom control commands with magic byte validation.")
    ]
    for i, (fn, desc) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(6)
        r1 = p.add_run()
        r1.text = "• " + fn + ": "
        r1.font.bold = True
        r1.font.size = Pt(12)
        r1.font.color.rgb = ACCENT_CYAN
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_WHITE

    add_card(slide4, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), title="Kernel Safety & Isolation", border_color=ACCENT_ORANGE)
    tb2 = slide4.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.3))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    safety = [
        ("Memory Boundary Isolation", "Kernel code NEVER dereferences `__user` pointers directly. `copy_from_user()` verifies page validity and handles page faults safely."),
        ("Mutex Synchronization", "Shared buffer and length variables are protected by `DEFINE_MUTEX(mydev_lock)`. Prevents race conditions from concurrent multi-process access."),
        ("Interruptible Waiting", "`mutex_lock_interruptible()` allows user processes waiting for the lock to be cleanly interrupted by signals (e.g. Ctrl+C)."),
        ("Strict Error Unwinding", "`mydev_init()` uses reverse-order cleanup labels (`goto err_...`) to release device numbers and classes if registration fails.")
    ]
    for i, (st, sd) in enumerate(safety):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.space_before = Pt(10)
        r1 = p.add_run()
        r1.text = "• " + st + ": "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = ACCENT_ORANGE
        r2 = p.add_run()
        r2.text = sd
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_WHITE

    # =========================================================================
    # SLIDE 5: ioctl Interface Design
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "I/O Control (`ioctl`) & Shared Protocol")

    add_card(slide5, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2), title="Why ioctl? Out-of-Band Control Protocol", border_color=ACCENT_CYAN)
    tb = slide5.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Standard read() and write() only transfer raw byte streams. ioctl() ('Input/Output Control') is used for device configuration, status querying, and administrative operations."
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_WHITE

    p2 = tf.add_paragraph()
    p2.text = "Command Encoding Architecture (`mydev_ioctl.h`):"
    p2.font.size = Pt(14)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_CYAN
    p2.space_before = Pt(14)

    commands = [
        ("MYDEV_IOC_CLEAR = _IO('m', 0)", "Erase Buffer: Takes no arguments. Driver zeroes `device_buffer` and resets data_length to 0."),
        ("MYDEV_IOC_GET_SIZE = _IOR('m', 1, int)", "Query Capacity: Kernel writes buffer capacity (1024 bytes) back to user int pointer via put_user()."),
        ("MYDEV_IOC_GET_LEN = _IOR('m', 2, int)", "Query Used Bytes: Kernel writes active data length (under mutex lock) to user pointer.")
    ]
    for cmd, desc in commands:
        pc = tf.add_paragraph()
        pc.space_before = Pt(10)
        r1 = pc.add_run()
        r1.text = "• " + cmd + "\n   "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = ACCENT_GREEN
        r2 = pc.add_run()
        r2.text = desc
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_MUTED

    p3 = tf.add_paragraph()
    p3.text = "• Magic Byte Protection: Command uses type 'm'. Unrecognized or foreign ioctls return -ENOTTY safely."
    p3.font.size = Pt(12)
    p3.font.color.rgb = ACCENT_ORANGE
    p3.space_before = Pt(12)

    # =========================================================================
    # SLIDE 6: C++ User-Space Application
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "C++17 User Application Design (`app.cpp`)")

    add_card(slide6, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2), title="RAII & Resource Safety (`MyDevice`)", border_color=ACCENT_BLUE)
    tb = slide6.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.3), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True
    raii_pts = [
        ("Constructor Acquisition", "Calls `::open('/dev/mydev', O_RDWR)`. Throws `std::system_error` with errno if driver is not loaded."),
        ("Destructor Release", "`~MyDevice()` invokes `::close(fd_)` automatically upon scope exit, preventing descriptor leaks."),
        ("Deleted Copy Operations", "`MyDevice(const MyDevice&) = delete` stops accidental copying that could cause double-close bugs."),
        ("Encapsulated Operations", "Provides clean high-level methods: `writeString()`, `readAll()`, `clear()`, `capacity()`, `used()`.")
    ]
    for title, desc in raii_pts:
        p = tf.paragraphs[0] if title == raii_pts[0][0] else tf.add_paragraph()
        p.space_before = Pt(8)
        r1 = p.add_run()
        r1.text = "• " + title + ": "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = ACCENT_CYAN
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_WHITE

    add_card(slide6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), title="User Experience & Menu Flow", border_color=ACCENT_GREEN)
    tb2 = slide6.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.3))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    menu_pts = [
        ("Option 1: Write String", "Overwrites buffer from offset 0, notifies if payload exceeds 1 KB buffer size."),
        ("Option 2: Read Buffer", "Reads accumulated text using rewind() + chunked read loop until EOF (0 bytes)."),
        ("Option 3: Clear Buffer", "Executes `MYDEV_IOC_CLEAR` ioctl to zero the driver memory."),
        ("Option 4: Query Size", "Executes `MYDEV_IOC_GET_SIZE` and `MYDEV_IOC_GET_LEN` ioctl calls to report memory metrics."),
        ("Robust Exception Catching", "Each menu operation is wrapped in a try/catch block so failed syscalls don't terminate the CLI.")
    ]
    for title, desc in menu_pts:
        p = tf2.paragraphs[0] if title == menu_pts[0][0] else tf2.add_paragraph()
        p.space_before = Pt(8)
        r1 = p.add_run()
        r1.text = "• " + title + ": "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = ACCENT_GREEN
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_WHITE

    # =========================================================================
    # SLIDE 7: Build System & Udev Permissions
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "Build System & Device Permissions")

    add_card(slide7, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2), title="Two-Pass Kbuild Makefile", border_color=ACCENT_CYAN)
    tb = slide7.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.3), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True
    make_pts = [
        ("Pass 1 (User Invocation)", "Runs when `KERNELRELEASE` is empty. Detects kernel build path via `KDIR ?= /lib/modules/$(shell uname -r)/build`."),
        ("Kbuild Delegation", "`make -C $(KDIR) M=$(CURDIR) modules` delegates compilation to kernel's official build system with exact compiler flags."),
        ("Pass 2 (Kbuild Invocation)", "Kbuild sets `KERNELRELEASE` and processes `obj-m := mydev.o` to produce `mydev.ko`."),
        ("User App Target", "Compiles `app.cpp` using `g++ -std=c++17 -Wall -Wextra -O2`.")
    ]
    for title, desc in make_pts:
        p = tf.paragraphs[0] if title == make_pts[0][0] else tf.add_paragraph()
        p.space_before = Pt(8)
        r1 = p.add_run()
        r1.text = "• " + title + ": "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = ACCENT_CYAN
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_WHITE

    add_card(slide7, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), title="Udev Permissions (`mydev.sh`)", border_color=ACCENT_ORANGE)
    tb2 = slide7.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.3))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    udev_pts = [
        ("The Permission Problem", "Default `/dev` device nodes created by the kernel are owned by root with mode 0600 (unprivileged `./app` gets 'Permission Denied')."),
        ("Udev Rule Installation", "`mydev.sh` installs `/etc/udev/rules.d/99-mydev.rules` with rule: `KERNEL==\"mydev\", MODE=\"0666\"`."),
        ("Automated Reload", "Reloads udev daemon rules (`udevadm control --reload-rules`) and awaits event settlement (`udevadm settle`)."),
        ("Clean Teardown", "On `unload`, cleans up the udev rule file so the host OS remains pristine.")
    ]
    for title, desc in udev_pts:
        p = tf2.paragraphs[0] if title == udev_pts[0][0] else tf2.add_paragraph()
        p.space_before = Pt(8)
        r1 = p.add_run()
        r1.text = "• " + title + ": "
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = ACCENT_ORANGE
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_WHITE

    # =========================================================================
    # SLIDE 8: Live Verification & Testing Results
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Live Verification & Test Results")

    add_card(slide8, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2), title="User App Verification (`./app`)", border_color=ACCENT_GREEN)
    tb = slide8.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.3), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True
    output_lines = [
        "Opened /dev/mydev - file descriptor = 3",
        "",
        "1) Write String -> 'Hello Linux Kernel World from C++!'",
        "   Result: Wrote 34 byte(s) to /dev/mydev",
        "",
        "2) Read Buffer",
        "   Result: Read 34 byte(s): \"Hello Linux Kernel World from C++!\"",
        "",
        "4) Query Buffer Size (ioctl)",
        "   Result: Capacity: 1024 bytes | In use: 34 bytes",
        "",
        "3) Clear Buffer (ioctl) -> Buffer cleared.",
        "2) Read Buffer -> Buffer is empty."
    ]
    for line in output_lines:
        p = tf.paragraphs[0] if line == output_lines[0] else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(11)
        p.font.name = "Consolas"
        p.font.color.rgb = ACCENT_GREEN if "Result" in line or "Opened" in line else TEXT_WHITE

    add_card(slide8, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), title="Kernel Space Log (`dmesg | grep mydev`)", border_color=ACCENT_CYAN)
    tb2 = slide8.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.3))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    dmesg_lines = [
        "[ 132.340041] mydev: loaded - /dev/mydev ready",
        "              (major 240, minor 0, buffer 1024 bytes)",
        "[ 159.572095] mydev: device opened (open handles: 1)",
        "[ 159.572262] mydev: wrote 34 bytes (holds 34/1024)",
        "[ 159.572273] mydev: read 34 bytes (new position 34)",
        "[ 159.572282] mydev: ioctl GET_SIZE -> 1024",
        "[ 159.572285] mydev: ioctl GET_LEN -> 34",
        "[ 159.572289] mydev: ioctl CLEAR - buffer erased",
        "[ 159.572296] mydev: device closed (open handles: 0)"
    ]
    for line in dmesg_lines:
        p = tf2.paragraphs[0] if line == dmesg_lines[0] else tf2.add_paragraph()
        p.text = line
        p.font.size = Pt(11)
        p.font.name = "Consolas"
        p.font.color.rgb = ACCENT_CYAN

    # =========================================================================
    # SLIDE 9: Course Concepts Mapping
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9)
    add_header(slide9, "Mapping to Core Operating System Concepts")

    quads = [
        (Inches(0.8), Inches(1.6), Inches(5.7), Inches(2.5),
         "1. Kernel vs User Space & Ring Isolation",
         "CPU Privilege Levels: Ring 3 (User) has restricted memory access. Ring 0 (Kernel) has direct hardware access. Kernel traps protect against invalid pointer dereferences via copy_to/from_user.",
         ACCENT_CYAN),
        (Inches(6.8), Inches(1.6), Inches(5.7), Inches(2.5),
         "2. File Descriptors & VFS Abstraction",
         "'Everything is a file': The VFS routes generic read/write calls to driver-specific file_operations. File descriptors maintain per-open file positions (f_pos) independent across processes.",
         ACCENT_BLUE),
        (Inches(0.8), Inches(4.3), Inches(5.7), Inches(2.5),
         "3. Concurrency & Synchronization",
         "Multi-core Kernel Preemption: Shared memory structures (device_buffer) require mutex locking (mydev_lock). Atomic variables (atomic_t) allow lockless counter updates.",
         ACCENT_ORANGE),
        (Inches(6.8), Inches(4.3), Inches(5.7), Inches(2.5),
         "4. Hardware Abstraction Layer (HAL)",
         "Modular Driver Architecture: The user application is completely agnostic to hardware details. The 1 KB RAM buffer can be replaced with physical UART/SPI without modifying app.cpp.",
         ACCENT_GREEN)
    ]
    for left, top, w, h, q_title, q_desc, col in quads:
        add_card(slide9, left, top, w, h, title=q_title, border_color=col)
        tb_q = slide9.shapes.add_textbox(left + Inches(0.2), top + Inches(0.6), w - Inches(0.4), h - Inches(0.7))
        tf_q = tb_q.text_frame
        tf_q.word_wrap = True
        pq = tf_q.paragraphs[0]
        pq.text = q_desc
        pq.font.size = Pt(12)
        pq.font.color.rgb = TEXT_WHITE

    # =========================================================================
    # SLIDE 10: Conclusion & Key Takeaways
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10)
    add_header(slide10, "Summary & Key Takeaways")

    add_card(slide10, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2), title="Project Accomplishments & Takeaways", border_color=ACCENT_CYAN)
    tb = slide10.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True

    takeaways = [
        ("End-to-End Systems Mastery: ", "Successfully designed, compiled, loaded, and verified a complete Linux character device driver and user-space client."),
        ("Safe Kernel Practices: ", "Mastered safe user-kernel memory boundaries (copy_from_user, put_user), dynamic device allocation, and mutex locking."),
        ("Modern C++ in Systems Programming: ", "Applied RAII to OS resources, guaranteeing deterministic file descriptor release without memory/resource leaks."),
        ("Extensibility to Physical Hardware: ", "The exact same driver architecture directly applies to physical embedded hardware (I2C sensors, SPI LCDs, GPIO pins, UART ports).")
    ]
    for i, (t_title, t_desc) in enumerate(takeaways):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(14)
        r1 = p.add_run()
        r1.text = "✔ " + t_title
        r1.font.bold = True
        r1.font.size = Pt(14)
        r1.font.color.rgb = ACCENT_CYAN
        r2 = p.add_run()
        r2.text = t_desc
        r2.font.size = Pt(13)
        r2.font.color.rgb = TEXT_WHITE

    prs.save(output_path)
    print(f"[+] Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    output_file = "/mnt/d/virtual-char-device-driver/virtual_char_device_driver.pptx"
    create_presentation(output_file)
